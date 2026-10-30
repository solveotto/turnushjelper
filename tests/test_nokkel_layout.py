"""The nøkkel grid must read exactly like the employer's Excel.

The Excel lays its 6 groups out column-major: each column, read top to
bottom, is six consecutive calendar weeks, and the next column carries on
where it stopped. A rutetermin rarely starts in group 1 — R27 starts in group
5 (14.12.26), R25 in group 2 — so the groups before the starting one have no
week in the first column and their labels begin one column to the right.

Until 2026-10-30 the readers sliced nine columns per group from that group's
own first "Uke" label (the R27 layout had been taken for a broken draft). The
turnusnøkkel page then showed groups 5-6 one week left of groups 1-4, while
mintur kept a fixed H-P: it lost R27's column Q (29.11.27-12.12.27) from the
page and the iCal export, and its week labels sat one column off. The old
tests built an artificial "shifted" file and only checked that dates appeared
somewhere in the HTML, so neither showed.

Every grid here — synthetic, read from the real files, rendered on the page,
built for mintur — goes through one calendar checker, and the page and
mintur must match the Excel cell for cell.
"""

import datetime
import json
import os
import re

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string
from sqlalchemy.orm import sessionmaker

from app.models import TurnusSet
from app.utils import kompdag_utils, turnuskalender as tk
from tests.conftest import login_user

#: The date grid's columns in each year's Excel, read off the files by hand.
EXCEL_COLUMNS = {"R25": ("H", "Q"), "R26": ("H", "P"), "R27": ("H", "Q")}
#: A turnus present in each year's schedule.
TURNUS = {"R25": "OSL_01", "R26": "OSL_01", "R27": "OSL_01_ALSTOM"}


def _needs(year):
    return pytest.mark.skipif(
        not os.path.exists(tk.nokkel_path(year)),
        reason=f"turnusdata/{year.lower()} is gitignored and not present",
    )


REAL_YEARS = [pytest.param(year, marks=_needs(year)) for year in EXCEL_COLUMNS]


# --- synthetic nøkkel -------------------------------------------------------


class _Cell:
    """Minimal stand-in for an openpyxl cell."""

    def __init__(self, value=None):
        self.value = value


def _blank_rows():
    return [[_Cell() for _ in range(22)] for _ in range(tk.GRID_ROWS)]


def _build_rows(first_monday, weeks, start_group, empty_labels=()):
    """A nøkkel laid out the way the employer's file is.

    Week ``k`` of the rutetermin sits in group ``(start_group + k) % 6``,
    column ``H + (start_group + k) // 6``. ``empty_labels`` adds
    ``(group, column, text)`` headers with no dates under them, as R27 has
    in Q17, Q25 and Q33.
    """
    rows = _blank_rows()
    for k in range(weeks):
        slot = start_group + k
        group, column = slot % 6, 7 + slot // 6
        monday = first_monday + datetime.timedelta(weeks=k)
        rows[group * 8][column] = _Cell(f"Uke {monday.isocalendar()[1]}")
        for d in range(7):
            day = monday + datetime.timedelta(days=d)
            rows[group * 8 + 1 + d][column] = _Cell(
                datetime.datetime(day.year, day.month, day.day)
            )
    for group, column, text in empty_labels:
        rows[group * 8][column] = _Cell(text)
    return rows


SHAPES = {
    "R25": dict(first_monday=datetime.date(2024, 12, 9), weeks=54, start_group=1),
    "R26": dict(first_monday=datetime.date(2025, 12, 15), weeks=54, start_group=0),
    "R27": dict(
        first_monday=datetime.date(2026, 12, 14),
        weeks=52,
        start_group=4,
        empty_labels=((2, 16, "Uke 50"), (3, 16, "Uke 51"), (4, 16, "Uke 52")),
    ),
}


# --- grids: {"labels": [...], "dates": 7 rows of one entry per column} -------


def _text(value):
    if value is None:
        return ""
    return value.strftime("%d.%m.%y") if hasattr(value, "strftime") else str(value)


def _grid_from_rows(rows):
    """The grid as the readers in turnuskalender see it, as text."""
    columns = tk.grid_columns(rows)
    return [
        {
            "labels": tk.week_labels(rows, g),
            "dates": [
                [_text(c.value) for c in rows[g * 8 + 1 + d][columns]]
                for d in range(7)
            ],
        }
        for g in range(6)
    ]


def _excel_grid(year):
    """The date grid as the Excel shows it, read straight from its cells."""
    first, last = (column_index_from_string(c) for c in EXCEL_COLUMNS[year])
    wb = openpyxl.load_workbook(tk.nokkel_path(year), data_only=True)
    try:
        ws = wb[tk.SHEET_NAME]
        # Nothing of the grid may lie outside the columns read by hand.
        for row in ws.iter_rows(min_row=1, max_row=tk.GRID_ROWS):
            for cell in row:
                is_label = str(cell.value or "").strip().lower().startswith("uke")
                if is_label or tk.cell_date(cell):
                    assert first <= cell.column <= last, cell.coordinate

        def text(row, column):
            return _text(ws.cell(row=row, column=column).value)

        columns = range(first, last + 1)
        return [
            {
                "labels": [text(g * 8 + 1, c) for c in columns],
                "dates": [[text(g * 8 + 2 + d, c) for c in columns] for d in range(7)],
            }
            for g in range(6)
        ]
    finally:
        wb.close()


def _page_grid(html):
    """The date grid of the rendered turnusnøkkel page, one table per group."""
    grid = []
    for table in html.split("<table>")[1:]:
        labels = re.findall(r'<th class="date-header">([^<]*)</th>', table)
        cells = re.findall(r'<td class="date-cell[^"]*">([^<]*)</td>', table)
        assert len(cells) == 7 * len(labels), "date cells don't fill the week columns"
        width = len(labels)
        grid.append(
            {"labels": labels, "dates": [cells[d * width:(d + 1) * width] for d in range(7)]}
        )
    assert len(grid) == 6
    return grid


def _mintur_grid(groups):
    return [
        {
            "labels": group["uke_labels"],
            "dates": [[d["value"] for d in row["dates"]] for row in group["day_rows"]],
        }
        for group in groups
    ]


def _assert_calendar_grid(grid):
    """The grid reads as a calendar, the way the Excel does.

    Every group is equally wide and has a label per column; each filled column
    of a group is one Monday-first week under its own "Uke NN"; and reading
    the columns left to right, each top to bottom, gives consecutive weeks
    with blanks only before the first and after the last.
    """
    width = len(grid[0]["labels"])
    weeks = []
    for g, group in enumerate(grid, start=1):
        assert len(group["labels"]) == width, f"group {g}: {len(group['labels'])} labels"
        assert len(group["dates"]) == 7
        for row in group["dates"]:
            assert len(row) == width, f"group {g}: {len(row)} date columns, {width} labels"
    for column in range(width):
        for g, group in enumerate(grid, start=1):
            texts = [group["dates"][d][column] for d in range(7)]
            if not any(texts):
                weeks.append(None)
                continue
            where = f"group {g}, column {column + 1}"
            assert all(texts), f"{where} is partly empty: {texts}"
            days = [datetime.datetime.strptime(t, "%d.%m.%y").date() for t in texts]
            monday = days[0]
            assert monday.weekday() == 0, f"{where} starts on {monday}"
            assert days == [monday + datetime.timedelta(days=d) for d in range(7)], (
                f"{where} is not one week: {texts}"
            )
            label = group["labels"][column].strip()
            assert label == f"Uke {monday.isocalendar()[1]}", (
                f"{where}: label {label!r} over the week of {monday}"
            )
            weeks.append(monday)
    filled = [i for i, monday in enumerate(weeks) if monday is not None]
    assert filled, "the grid holds no dates"
    run = weeks[filled[0]:filled[-1] + 1]
    assert None not in run, "a blank week inside the grid"
    for before, after in zip(run, run[1:]):
        assert after - before == datetime.timedelta(weeks=1), (
            f"the week of {after} follows the week of {before} down the columns"
        )


# --- the column span ---------------------------------------------------------


class TestGridColumns:
    """One span for the whole grid, from its leftmost to its rightmost week."""

    def test_r26_shape_is_h_to_p(self):
        assert tk.grid_columns(_build_rows(**SHAPES["R26"])) == slice(7, 16)

    def test_r27_shape_is_h_to_q(self):
        assert tk.grid_columns(_build_rows(**SHAPES["R27"])) == slice(7, 17)

    def test_r25_shape_is_h_to_q(self):
        assert tk.grid_columns(_build_rows(**SHAPES["R25"])) == slice(7, 17)

    def test_label_match_is_case_insensitive_and_tolerates_padding(self):
        rows = _blank_rows()
        rows[0][9] = _Cell("  uke 53 ")
        assert tk.grid_columns(rows) == slice(9, 10)

    def test_falls_back_to_h_to_p_without_labels_or_dates(self):
        assert tk.grid_columns(_blank_rows()) == tk.DATE_COLUMNS


@pytest.mark.parametrize("year", SHAPES)
def test_synthetic_grid_reads_as_a_calendar(year):
    _assert_calendar_grid(_grid_from_rows(_build_rows(**SHAPES[year])))


@pytest.mark.parametrize("year", SHAPES)
def test_synthetic_index_has_every_day_on_its_weekday(year):
    shape = SHAPES[year]
    index = tk.build_date_index(rows=_build_rows(**shape))
    assert len(index) == shape["weeks"] * 7
    assert all(day.weekday() == weekday for day, (_, weekday) in index.items())


# --- the real files ----------------------------------------------------------


@pytest.mark.parametrize("year", REAL_YEARS)
def test_real_grid_is_the_excel_grid(year):
    grid = _grid_from_rows(tk.load_nokkel_rows(year))
    assert grid == _excel_grid(year)
    _assert_calendar_grid(grid)


@_needs("R27")
def test_r27_covers_its_rutetermin():
    index = tk.build_date_index("R27")
    assert min(index) == datetime.date(2026, 12, 14)
    assert max(index) == datetime.date(2027, 12, 12)
    assert len(index) == 364
    assert tk.grid_columns(tk.load_nokkel_rows("R27")) == slice(7, 17)


@pytest.mark.parametrize("year", REAL_YEARS)
def test_holiday_positions_match_the_sheet(year):
    """Every kompdag holiday is found, at the group and weekday the sheet
    puts it — checked against a scan of the whole sheet, not the span."""
    rows = tk.load_nokkel_rows(year)
    where = {}
    for g in range(6):
        for d in range(7):
            for cell in rows[g * 8 + 1 + d]:
                day = tk.cell_date(cell)
                if day is not None:
                    where[day] = (g, d)

    positions = kompdag_utils.get_holiday_positions(year)

    assert {day for _, _, day in positions} == kompdag_utils.get_kompdag_holidays(where)
    assert all(where[day] == (g, d) for g, d, day in positions)


# --- what the driver sees ----------------------------------------------------


def _seed(db_session, year):
    ts = TurnusSet(name=year, year_identifier=year, is_active=1)
    db_session.add(ts)
    db_session.commit()
    return ts


def _place_driver(monkeypatch, ts, year, linje=1):
    record = {
        "shift_title": TURNUS[year],
        "linjenummer": linje,
        "turnus_set_id": ts.id,
        "year_identifier": year,
    }
    monkeypatch.setattr(
        "app.routes.shifts.mintur.get_innplassering_for_user", lambda uid: [record]
    )


@pytest.mark.parametrize("year", REAL_YEARS)
def test_turnusnokkel_page_is_the_excel_grid(
    year, client, sample_user, db_session, monkeypatch
):
    ts = _seed(db_session, year)
    TestSession = sessionmaker(bind=db_session.get_bind())
    monkeypatch.setattr(
        "app.routes.shifts.turnusnokkel.get_db_session", lambda: TestSession()
    )
    login_user(client, sample_user["username"], sample_user["password"])

    resp = client.get(f"/turnusnokkel/{ts.id}/{TURNUS[year]}")

    assert resp.status_code == 200
    html = resp.get_data(as_text=True)
    grid = _page_grid(html)
    assert grid == _excel_grid(year)
    _assert_calendar_grid(grid)

    # Each group is its own <table>. With automatic widths an empty week
    # column collapsed to a sliver and the groups no longer lined up
    # (2026-10-30), so every table declares the same fixed columns, one per
    # week column — blank ones included.
    colgroups = re.findall(r"<colgroup>(.*?)</colgroup>", html, re.S)
    assert len(colgroups) == 6
    assert len(set(colgroups)) == 1, "the groups' tables declare different columns"
    assert colgroups[0].count('<col class="col-date">') == len(grid[0]["labels"])


def _excel_red_dates(year):
    """Dates the Excel prints in red font — its own holiday marking."""
    wb = openpyxl.load_workbook(tk.nokkel_path(year))
    try:
        return {
            _text(cell.value)
            for row in wb[tk.SHEET_NAME].iter_rows(min_row=1, max_row=tk.GRID_ROWS)
            for cell in row
            if tk.cell_date(cell) is not None
            and cell.font.color is not None
            and cell.font.color.type == "rgb"
            and str(cell.font.color.rgb).upper().endswith("FF0000")
        }
    finally:
        wb.close()


@pytest.mark.parametrize("year", REAL_YEARS)
def test_page_marks_every_date_the_excel_marks_red(
    year, client, sample_user, db_session, monkeypatch
):
    """The page marks the computed §5.13.1 set, not the red font. That set
    must still cover every date the Excel prints red; it may add more, since
    the Excel misses påskeaften and 1. pinsedag."""
    ts = _seed(db_session, year)
    TestSession = sessionmaker(bind=db_session.get_bind())
    monkeypatch.setattr(
        "app.routes.shifts.turnusnokkel.get_db_session", lambda: TestSession()
    )
    login_user(client, sample_user["username"], sample_user["password"])

    html = client.get(f"/turnusnokkel/{ts.id}/{TURNUS[year]}").get_data(as_text=True)

    marked = set(re.findall(r'<td class="date-cell holiday">([^<]*)</td>', html))
    red = _excel_red_dates(year)
    assert red, "the Excel marks no date red"
    assert red <= marked, f"red in the Excel, not on the page: {sorted(red - marked)}"


@pytest.mark.parametrize("year", REAL_YEARS)
def test_mintur_grid_is_the_excel_grid(year, patch_db, db_session, monkeypatch):
    from app.routes.shifts.mintur import _load_mintur_data

    ts = _seed(db_session, year)
    _place_driver(monkeypatch, ts, year)

    grid = _mintur_grid(_load_mintur_data(1)["groups"])

    assert grid == _excel_grid(year)
    _assert_calendar_grid(grid)


@_needs("R27")
def test_ical_export_keeps_the_last_weeks_of_r27(
    client, sample_user, db_session, monkeypatch
):
    """Column Q holds 29.11.27-12.12.27; the export lost it with mintur."""
    ts = _seed(db_session, "R27")
    _place_driver(monkeypatch, ts, "R27", linje=1)
    with open(tk.schedule_path("R27"), encoding="utf-8") as f:
        weeks = next(e["OSL_01_ALSTOM"] for e in json.load(f) if "OSL_01_ALSTOM" in e)
    index = tk.build_date_index("R27")
    worked = []
    for i in range(14):
        day = datetime.date(2027, 11, 29) + datetime.timedelta(days=i)
        group, weekday = index[day]
        cell = tk.lookup_cell(weeks, tk.rotation_week(group, 1), weekday)
        if cell and len(cell.get("tid", [])) >= 2:
            worked.append(day)
    assert worked, "linje 1 works no day in the last two weeks"
    login_user(client, sample_user["username"], sample_user["password"])

    body = client.get(f"/mintur/export_ical?sett={ts.id}").get_data(as_text=True)

    for day in worked:
        assert f"UID:OSL_01_ALSTOM-{day:%Y%m%d}@turnushjelper" in body, day
