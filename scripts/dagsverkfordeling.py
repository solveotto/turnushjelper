#!/usr/bin/env python3
"""Build the "Dagsverkfordeling <year> Oslo" sheet as PDF and/or Excel.

Mirrors the R26 sheet the union circulated: per turnus, how many duties start
before 06, 06:00-07:59, 08:00-11:59, how many are kveld and natt, the sum, and
how many of them are dobbelturer. The time categories come from the app's own
classifier (turnus_stats_<YEAR>.json, see app/utils/shift_stats.py), the
dobbelturer from double_shifts_<year>.json.

Usage:
    venv/bin/python scripts/dagsverkfordeling.py --year R27
    venv/bin/python scripts/dagsverkfordeling.py --year R27 --format xlsx --out-dir ~/Downloads
    venv/bin/python scripts/dagsverkfordeling.py --year R26 --label-1-8 "70 / 73" --label-16-19 Gjøvik

Group labels for turnus 1-8 and 16-19 default to the suffix in the turnus
names (ALSTOM, Gjøvik, ...); the other groups are fixed.
"""

import argparse
import json
import os
import re
import sys
from collections import Counter

project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, project_root)

STAT_COLS = ["before_6", "tidlig_6_8", "tidlig_8_12", "ettermiddag", "natt", "shift_cnt"]
HEADERS = ["Før 06", "06:00 - 07:59", "08:00 - 11:59", "Etter 12 (Kveld)", "Natt",
           "Sum dagsverk. (Inkl H)", "Herav dobbelturer"]
FIXED_GROUPS = [(9, 15, "Østre Linje"), (20, 40, "Øvrige")]


# --- data ---------------------------------------------------------------

def row_label(name):
    m = re.match(r"OSL_(\d+)", name)
    if m:
        n = int(m.group(1))
        return "40 (3. helg)" if n == 40 else str(n)
    m = re.match(r"OSL_Ramme_(\d+)", name)
    if m:
        return f"Ramme {int(m.group(1))}"
    m = re.match(r"OSL_Utland_(\d+)", name)
    if m:
        return f"Utland {int(m.group(1))}"
    return name


def turnus_number(name):
    m = re.match(r"OSL_(\d+)", name)
    return int(m.group(1)) if m else None


def group_of(name, groups):
    n = turnus_number(name)
    if n is not None:
        for lo, hi, label in groups:
            if lo <= n <= hi:
                return label
    if "Ramme" in name:
        return "Rammer"
    if "Utland" in name:
        return "Utland"
    return ""


def label_from_names(names, lo, hi):
    """Most common name suffix in the range, e.g. OSL_16_ALSTOM -> ALSTOM."""
    suffixes = Counter()
    for name in names:
        n = turnus_number(name)
        if n is None or not lo <= n <= hi:
            continue
        m = re.match(r"OSL_\d+_(.+)", name)
        if m:
            suffixes[m.group(1).replace("_", " ")] += 1
    return suffixes.most_common(1)[0][0] if suffixes else f"{lo}-{hi}"


def base_number(dagsverk):
    m = re.match(r"\d+", str(dagsverk))
    return m.group(0) if m else ""


def rotation_days(week_dict):
    """All day dicts in rotation order (week 1 day 1 ... week 6 day 7)."""
    out = []
    for uke in sorted((k for k in week_dict if k.isdigit()), key=int):
        days = week_dict[uke]
        if not isinstance(days, dict):
            continue
        for dn in sorted((k for k in days if k.isdigit()), key=int):
            if isinstance(days[dn], dict):
                out.append(days[dn])
    return out


def count_dobbelturer(week_dict, first_to_second, second_to_first):
    """One per first leg, plus second legs whose first leg is not the day before.

    A second leg right after its first leg is the same dobbeltur and is not
    counted twice. A lone second leg (first leg printed under another number,
    or as '?') still counts. Reproduces the R26 sheet in every row.
    """
    days = rotation_days(week_dict)
    count = 0
    for i, d in enumerate(days):
        b = base_number(d.get("dagsverk"))
        if b in first_to_second:
            count += 1
        elif b in second_to_first:
            prev = base_number(days[i - 1].get("dagsverk"))
            if prev != second_to_first[b]:
                count += 1
    return count


def build_rows(year, data_dir, label_1_8=None, label_16_19=None):
    """Return [(group, label, [før6, 6-8, 8-12, kveld, natt, sum, dobbelturer]), ...]."""
    y = year.lower()
    with open(os.path.join(data_dir, f"turnus_stats_{year}.json"), encoding="utf-8") as f:
        stats = json.load(f)
    schedule = {}
    with open(os.path.join(data_dir, f"turnus_schedule_{year}.json"), encoding="utf-8") as f:
        for item in json.load(f):
            schedule.update(item)
    with open(os.path.join(data_dir, f"double_shifts_{y}.json"), encoding="utf-8") as f:
        dbl = json.load(f)
    first_to_second = {p["first_shift"].rstrip("-"): p["second_shift"].rstrip("-")
                       for p in dbl["dobbelt_tur"]}
    second_to_first = {v: k for k, v in first_to_second.items()}

    names = list(stats["turnus"].values())
    groups = [
        (1, 8, label_1_8 or label_from_names(names, 1, 8)),
        (16, 19, label_16_19 or label_from_names(names, 16, 19)),
    ] + FIXED_GROUPS

    rows = []
    for idx, name in stats["turnus"].items():
        nums = [stats[c][idx] for c in STAT_COLS]
        nums.append(count_dobbelturer(schedule[name], first_to_second, second_to_first))
        rows.append((group_of(name, groups), row_label(name), nums))
    return rows


def group_spans(rows, first_row):
    """Contiguous runs of the same group: [(label, first_row, last_row), ...]."""
    spans = []
    for i, (g, _label, _nums) in enumerate(rows):
        r = first_row + i
        if spans and spans[-1][0] == g:
            spans[-1][2] = r
        else:
            spans.append([g, r, r])
    return spans


# --- PDF ----------------------------------------------------------------

def build_pdf(year, rows, out_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = "Helvetica-Bold"

    class RotatedText(Flowable):
        """Text rotated 90 degrees, one or more lines, centred in its cell."""

        def __init__(self, lines, size=7):
            super().__init__()
            self.lines = lines
            self.size = size
            self.width = len(lines) * size * 1.25
            self.height = max(stringWidth(line, font, size) for line in lines)

        def draw(self):
            c = self.canv
            c.saveState()
            c.setFont(font, self.size)
            c.rotate(90)
            line_h = self.size * 1.25
            x_offset = self.width - line_h * 0.3
            for i, line in enumerate(self.lines):
                slot = len(self.lines) - 1 - i  # first line leftmost
                c.drawCentredString(self.height / 2, -(x_offset - slot * line_h) + line_h * 0.55, line)
            c.restoreState()

    header_lines = [["Sum dagsverk.", "(Inkl H)"] if "(Inkl H)" in h else [h] for h in HEADERS]
    table_data = [["", ""] + [RotatedText(lines) for lines in header_lines]]
    spans = group_spans(rows, first_row=1)
    span_starts = {s[1] for s in spans}
    for i, (g, label, nums) in enumerate(rows):
        cells = ["" if n == 0 else str(n) for n in nums]
        table_data.append([g if (i + 1) in span_starts else "", label] + cells)

    n_rows = len(table_data)
    t = Table(table_data, colWidths=[67, 65] + [48] * 7, rowHeights=[88] + [11.0] * (n_rows - 1))
    thin, thick = 0.5, 1.5
    style = [
        ("FONT", (0, 0), (-1, -1), font, 7),
        ("FONT", (0, 1), (0, -1), font, 8),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("GRID", (1, 1), (-1, -1), thin, colors.black),
        ("BOX", (0, 0), (-1, -1), thick, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), thick, colors.black),
        ("LINEAFTER", (0, 0), (0, -1), thick, colors.black),
        ("LINEAFTER", (1, 0), (1, -1), thick, colors.black),
        ("LINEBEFORE", (7, 0), (7, -1), thick, colors.black),
        ("LINEAFTER", (7, 0), (7, -1), thick, colors.black),
        ("LINEBEFORE", (2, 0), (-1, 0), thin, colors.black),
    ]
    for _g, r0, r1 in spans:
        style.append(("SPAN", (0, r0), (0, r1)))
        style.append(("LINEABOVE", (0, r0), (-1, r0), thick, colors.black))
    t.setStyle(TableStyle(style))

    doc = SimpleDocTemplate(
        out_path, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=10 * mm, bottomMargin=8 * mm,
        title=f"Dagsverkfordeling {year} Oslo",
    )
    title = Paragraph(f"Dagsverkfordeling {year} Oslo",
                      ParagraphStyle("t", fontName=font, fontSize=10, leading=12))
    doc.build([title, Spacer(1, 4), t])


# --- Excel --------------------------------------------------------------

def build_xlsx(year, rows, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = f"Dagsverkfordeling {year}"
    thin = Side(style="thin")
    thick = Side(style="medium")
    bold = Font(name="Arial", bold=True, size=8)
    center = Alignment(horizontal="center", vertical="center")

    def set_sides(cell, **sides):
        b = cell.border
        kw = dict(left=b.left, right=b.right, top=b.top, bottom=b.bottom)
        kw.update(sides)
        cell.border = Border(**kw)

    ws["A1"] = f"Dagsverkfordeling {year} Oslo"
    ws["A1"].font = Font(name="Arial", bold=True, size=11)

    hdr_row = 3
    for i, h in enumerate(HEADERS):
        c = ws.cell(row=hdr_row, column=3 + i, value=h)
        c.font = bold
        c.alignment = Alignment(horizontal="center", vertical="center", textRotation=90, wrap_text=True)
    ws.row_dimensions[hdr_row].height = 80

    first_data = hdr_row + 1
    spans = group_spans(rows, first_data)
    span_starts = {s[1]: s[0] for s in spans}
    for i, (_g, label, nums) in enumerate(rows):
        r = first_data + i
        ws.cell(row=r, column=1, value=span_starts.get(r))
        ws.cell(row=r, column=2, value=label)
        for j, n in enumerate(nums[:5]):
            ws.cell(row=r, column=3 + j, value=n)
        ws.cell(row=r, column=8, value=f"=SUM(C{r}:G{r})")
        ws.cell(row=r, column=9, value=nums[6])
        for col in range(1, 10):
            c = ws.cell(row=r, column=col)
            c.font = bold
            c.alignment = center
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if col >= 3:
                c.number_format = "0;-0;;@"  # hide zeros like the original print
    last_data = first_data + len(rows) - 1

    for _g, s, e in spans:
        ws.merge_cells(start_row=s, start_column=1, end_row=e, end_column=1)
        ws.cell(row=s, column=1).alignment = center
        for col in range(1, 10):
            set_sides(ws.cell(row=s, column=col), top=thick)
    for row in range(hdr_row, last_data + 1):
        set_sides(ws.cell(row=row, column=1), left=thick, right=thick)
        set_sides(ws.cell(row=row, column=2), right=thick)
        set_sides(ws.cell(row=row, column=8), left=thick, right=thick)
        set_sides(ws.cell(row=row, column=9), right=thick)
    for col in range(1, 10):
        set_sides(ws.cell(row=hdr_row, column=col), top=thick, bottom=thick)
        set_sides(ws.cell(row=last_data, column=col), bottom=thick)

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 13
    for col in range(3, 10):
        ws.column_dimensions[get_column_letter(col)].width = 9

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:I{last_data}"
    wb.save(out_path)


# --- CLI ----------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Build the Dagsverkfordeling sheet for one rutetermin")
    parser.add_argument("--year", required=True, help="Year identifier, e.g. R27")
    parser.add_argument("--format", choices=["pdf", "xlsx", "both"], default="both")
    parser.add_argument("--out-dir", help="Output directory. Default: turnusdata/<year>/")
    parser.add_argument("--label-1-8", help="Group label for turnus 1-8 (default: suffix in turnus names)")
    parser.add_argument("--label-16-19", help="Group label for turnus 16-19 (default: suffix in turnus names)")
    args = parser.parse_args()

    from config import AppConfig

    year = args.year.upper()
    data_dir = os.path.join(AppConfig.turnusfiler_dir, year.lower())
    if not os.path.isdir(data_dir):
        print(f"Fant ikke {data_dir}", file=sys.stderr)
        return 1
    out_dir = os.path.expanduser(args.out_dir) if args.out_dir else data_dir
    os.makedirs(out_dir, exist_ok=True)

    rows = build_rows(year, data_dir, args.label_1_8, args.label_16_19)
    stem = os.path.join(out_dir, f"Dagsverkfordeling {year} Oslo")
    if args.format in ("pdf", "both"):
        build_pdf(year, rows, stem + ".pdf")
        print(f"Skrev {stem}.pdf")
    if args.format in ("xlsx", "both"):
        build_xlsx(year, rows, stem + ".xlsx")
        print(f"Skrev {stem}.xlsx")
    for g, label, nums in rows:
        print(f"{g:12} {label:14} " + " ".join(f"{n:3}" for n in nums))
    return 0


if __name__ == "__main__":
    sys.exit(main())
